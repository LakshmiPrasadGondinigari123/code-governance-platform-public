"""Evidence graph queries and exports (SARIF, CycloneDX, CI gate, Jira, webhook, Excel)."""
import io
import os
import uuid
from collections import deque
from typing import Optional
from urllib.parse import quote

import httpx
from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import JSONResponse, Response

from . import auth, db, scanner
from .common import SEV_ORDER, audit, findings_for, latest_scans, now, paginate, scan_or_404

router = APIRouter()


def _graph(user, scan_id):
    scan_or_404(user, scan_id)
    g = db.get("graphs", user["tenant_id"], scan_id)
    if not g:
        raise HTTPException(404, "graph not available (scan not completed?)")
    return g


def _adj(g, directed=False):
    adj = {}
    for e in g["edges"]:
        adj.setdefault(e["src"], []).append((e["dst"], e))
        if not directed:
            adj.setdefault(e["dst"], []).append((e["src"], e))
    return adj


def _node(g, key):
    n = next((n for n in g["nodes"] if n["key"] == key), None)
    if not n:
        raise HTTPException(404, "node not found")
    return n


# ------------------------------------------------------------------ graph
@router.get("/scans/{scan_id}/graph/nodes")
def graph_nodes(scan_id: str, type: Optional[str] = None, limit: int = 200, offset: int = 0, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    rows = [n for n in g["nodes"] if not type or n["type"] == type]
    return {**paginate(rows, limit, offset), "types": sorted({n["type"] for n in g["nodes"]})}


@router.get("/scans/{scan_id}/graph/edges")
def graph_edges(scan_id: str, type: Optional[str] = None, limit: int = 500, offset: int = 0, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    return paginate([e for e in g["edges"] if not type or e["type"] == type], limit, offset)


@router.get("/scans/{scan_id}/graph/search")
def graph_search(scan_id: str, q: str = Query(..., min_length=1), limit: int = 50, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    ql = q.lower()
    return [n for n in g["nodes"] if ql in n["label"].lower() or ql in n["key"].lower()][: max(1, min(limit, 200))]


@router.get("/scans/{scan_id}/graph/path")
def graph_path(scan_id: str, source: str, target: str, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    _node(g, source), _node(g, target)
    adj = _adj(g)
    prev, q = {source: None}, deque([source])
    while q:
        cur = q.popleft()
        if cur == target:
            break
        for nxt, e in adj.get(cur, []):
            if nxt not in prev:
                prev[nxt] = (cur, e)
                q.append(nxt)
    if target not in prev:
        return {"found": False, "nodes": [], "edges": []}
    path, edges, cur = [], [], target
    while cur:
        path.append(cur)
        step = prev[cur]
        if step:
            edges.append(step[1])
            cur = step[0]
        else:
            cur = None
    by_key = {n["key"]: n for n in g["nodes"]}
    return {"found": True, "nodes": [by_key[k] for k in reversed(path)], "edges": list(reversed(edges))}


@router.get("/scans/{scan_id}/graph/nodes/{node_key:path}")
def graph_node_detail(scan_id: str, node_key: str, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    n = _node(g, node_key)
    inc = [e for e in g["edges"] if e["dst"] == node_key]
    out = [e for e in g["edges"] if e["src"] == node_key]
    return {"node": n, "incoming": inc, "outgoing": out, "degree": len(inc) + len(out)}


@router.get("/scans/{scan_id}/graph/neighborhood/{node_key:path}")
def graph_neighborhood(scan_id: str, node_key: str, depth: int = 1, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    _node(g, node_key)
    adj, depth = _adj(g), max(1, min(depth, 4))
    seen, frontier, edges = {node_key}, [node_key], {}
    for _ in range(depth):
        nxt_frontier = []
        for k in frontier:
            for nb, e in adj.get(k, []):
                edges[e["id"]] = e
                if nb not in seen:
                    seen.add(nb)
                    nxt_frontier.append(nb)
        frontier = nxt_frontier[:500]
    by_key = {n["key"]: n for n in g["nodes"]}
    return {"center": node_key, "nodes": [by_key[k] for k in seen if k in by_key], "edges": list(edges.values())}


@router.get("/scans/{scan_id}/graph/finding-review-status")
def graph_review_status(scan_id: str, user: dict = Depends(auth.current_user)):
    scan_or_404(user, scan_id)
    out = {}
    for f in findings_for(user["tenant_id"], {scan_id}):
        out[f"finding:{f['id']}"] = {"classification": f.get("classification", "open"), "governance_status": f.get("governance_status", "unreviewed")}
    return out


@router.get("/scans/{scan_id}/graph/impact/{node_key:path}")
def graph_impact(scan_id: str, node_key: str, direction: str = "both", depth: int = 3, user: dict = Depends(auth.current_user)):
    """Everything that would be affected if this node (package/model/provider/host) changed."""
    g = _graph(user, scan_id)
    _node(g, node_key)
    rev = {}
    for e in g["edges"]:
        rev.setdefault(e["dst"], []).append(e["src"])
    seen, frontier = {node_key}, [node_key]
    for _ in range(max(1, min(depth, 6))):
        nxt = []
        for k in frontier:
            for s in rev.get(k, []):
                if s not in seen:
                    seen.add(s)
                    nxt.append(s)
        frontier = nxt
    by_key = {n["key"]: n for n in g["nodes"]}
    impacted = [by_key[k] for k in seen if k != node_key and k in by_key]
    by_type = {}
    for n in impacted:
        by_type[n["type"]] = by_type.get(n["type"], 0) + 1
    return {"node": node_key, "impacted": impacted, "by_type": by_type}


@router.get("/scans/{scan_id}/graph/edges/lookup")
def graph_edge_lookup(scan_id: str, source: str, target: str, type: Optional[str] = None, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    rows = [e for e in g["edges"] if e["src"] == source and e["dst"] == target and (not type or e["type"] == type)]
    if not rows:
        raise HTTPException(404, "edge not found")
    return rows


@router.get("/scans/{scan_id}/graph/edges/{edge_id}")
def graph_edge_detail(scan_id: str, edge_id: str, user: dict = Depends(auth.current_user)):
    g = _graph(user, scan_id)
    e = next((e for e in g["edges"] if e["id"] == edge_id), None)
    if not e:
        raise HTTPException(404, "edge not found")
    by_key = {n["key"]: n for n in g["nodes"]}
    return {"edge": e, "source": by_key.get(e["src"]), "target": by_key.get(e["dst"])}


# ------------------------------------------------------------------ exports
@router.get("/scans/{scan_id}/export/sarif")
def export_sarif(scan_id: str, user: dict = Depends(auth.current_user)):
    s = scan_or_404(user, scan_id)
    fs = [f for f in findings_for(user["tenant_id"], {scan_id}) if f["category"] != "ai_governance"]
    level = {"critical": "error", "high": "error", "medium": "warning", "low": "note", "info": "note"}
    rules = {f["rule_id"]: {"id": f["rule_id"], "name": f["title"], "shortDescription": {"text": f["title"]}} for f in fs}
    sarif = {
        "$schema": "https://json.schemastore.org/sarif-2.1.0.json", "version": "2.1.0",
        "runs": [{
            "tool": {"driver": {"name": "Code-to-Governance Intelligence Platform", "version": "1.0.0", "rules": list(rules.values())}},
            "results": [{
                "ruleId": f["rule_id"], "level": level[f["severity"]], "message": {"text": f"{f['title']}: {f['description']}"},
                "locations": [{"physicalLocation": {"artifactLocation": {"uri": f["file"]}, "region": {"startLine": max(1, f["line"])}}}],
                "partialFingerprints": {"primaryLocationLineHash": f["fingerprint"]},
            } for f in fs],
        }],
    }
    audit(user, "export.sarif", scan_id)
    return JSONResponse(sarif, headers={"Content-Disposition": f'attachment; filename="{s["repository_name"]}-sarif.json"'})


@router.get("/scans/{scan_id}/export/cyclonedx")
def export_cyclonedx(scan_id: str, user: dict = Depends(auth.current_user)):
    s = scan_or_404(user, scan_id)
    inv = db.get("inventory", user["tenant_id"], scan_id) or {"components": [], "models": [], "ai_components": []}
    comps = []
    for c in inv["components"]:
        eco = "pypi" if c["ecosystem"] == "pypi" else "npm"
        ver = c.get("version") or ""
        comps.append({"type": "library", "name": c["name"], "version": ver,
                      "purl": f"pkg:{eco}/{quote(c['name'])}" + (f"@{quote(ver)}" if ver else "")})
    for m in inv["models"]:
        comps.append({"type": "machine-learning-model", "name": m["name"], "publisher": m["provider"]})
    bom = {"bomFormat": "CycloneDX", "specVersion": "1.5", "version": 1, "serialNumber": f"urn:uuid:{uuid.uuid5(uuid.NAMESPACE_URL, scan_id)}",
           "metadata": {"timestamp": now(), "component": {"type": "application", "name": s["repository_name"]},
                        "tools": [{"name": "Code-to-Governance Intelligence Platform", "version": "1.0.0"}]},
           "components": comps}
    audit(user, "export.cyclonedx", scan_id)
    return JSONResponse(bom, headers={"Content-Disposition": f'attachment; filename="{s["repository_name"]}-cyclonedx.json"'})


@router.get("/scans/{scan_id}/ci-gate")
def ci_gate(scan_id: str, fail_on: str = "high", strict: bool = False, user: dict = Depends(auth.current_user)):
    s = scan_or_404(user, scan_id)
    if fail_on not in SEV_ORDER:
        raise HTTPException(400, "fail_on must be critical, high, medium, low or info")
    if s["status"] != "completed":
        raise HTTPException(409, f"scan is {s['status']}")
    fs = findings_for(user["tenant_id"], {scan_id})
    blocking = [f for f in fs if f["category"] != "ai_governance" and SEV_ORDER[f["severity"]] >= SEV_ORDER[fail_on]
                and f.get("classification", "open") == "open"]
    violated = [c["control_id"] for c in s.get("controls", []) if c["status"] == "violated"]
    passed = not blocking and not violated
    body = {"scan_id": scan_id, "passed": passed, "fail_on": fail_on, "blocking_findings": len(blocking), "violated_controls": violated,
            "blocking": [{"id": f["id"], "severity": f["severity"], "title": f["title"], "file": f["file"], "line": f["line"]} for f in blocking[:50]]}
    return JSONResponse(body, status_code=409 if (strict and not passed) else 200, headers={"X-Gate": "pass" if passed else "fail"})


def _jira_summary(s, fs):
    return [f for f in fs if f["severity"] in ("critical", "high") and f["category"] != "ai_governance" and f.get("classification") == "open"][:25]


@router.post("/scans/{scan_id}/export/jira")
def export_jira(scan_id: str, user: dict = Depends(auth.require("admin", "reviewer"))):
    s = scan_or_404(user, scan_id)
    todo = _jira_summary(s, findings_for(user["tenant_id"], {scan_id}))
    base, email, token, project = (os.getenv(k) for k in ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN", "JIRA_PROJECT_KEY"))
    payload = [{"summary": f"[{f['severity'].upper()}] {f['title']} in {s['repository_name']}",
                "description": f"{f['description']}\nFile: {f['file']}:{f['line']}\nRule: {f['rule_id']}"} for f in todo]
    if not all((base, email, token, project)):
        return {"status": "not_configured", "detail": "Set JIRA_BASE_URL, JIRA_EMAIL, JIRA_API_TOKEN and JIRA_PROJECT_KEY.", "issues_prepared": payload}
    created = []
    for p in payload:
        r = httpx.post(f"{base.rstrip('/')}/rest/api/2/issue", auth=(email, token), timeout=20, json={"fields": {
            "project": {"key": project}, "summary": p["summary"][:250], "description": p["description"], "issuetype": {"name": "Task"}}})
        created.append({"summary": p["summary"], "ok": r.status_code < 300, "key": (r.json().get("key") if r.status_code < 300 else None)})
    audit(user, "export.jira", scan_id, {"created": sum(1 for c in created if c["ok"])})
    return {"status": "sent", "issues": created}


@router.post("/scans/{scan_id}/export/webhook")
def export_webhook(scan_id: str, body: dict = Body(default={}), user: dict = Depends(auth.require("admin", "reviewer"))):
    s = scan_or_404(user, scan_id)
    url = body.get("url") or os.getenv("WEBHOOK_EXPORT_URL")
    if not url:
        return {"status": "not_configured", "detail": "Pass {\"url\": ...} or set WEBHOOK_EXPORT_URL."}
    try:
        scanner.validate_git_url(url)  # same SSRF guard: https/http, public hosts only
    except ValueError as e:
        raise HTTPException(400, str(e))
    payload = {"event": "scan.completed", "scan_id": scan_id, "repository": s["repository_name"], "summary": s.get("summary"), "sent_at": now()}
    r = httpx.post(url, json=payload, timeout=15, follow_redirects=False)
    audit(user, "export.webhook", scan_id, {"status": r.status_code})
    return {"status": "sent", "response_status": r.status_code}


# ------------------------------------------------------------------ Excel
def _workbook(scans: list, tenant: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2937")

    def sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Summary"
    sheet(ws, ["Repository", "Scan", "Finished", "Files", "Findings", "Critical", "High", "Medium", "Low", "Controls satisfied", "Controls total"],
          [[s["repository_name"], s["id"], s["finished_at"], s["summary"]["files_scanned"], s["summary"]["findings_total"],
            s["summary"]["by_severity"].get("critical", 0), s["summary"]["by_severity"].get("high", 0),
            s["summary"]["by_severity"].get("medium", 0), s["summary"]["by_severity"].get("low", 0),
            s["summary"]["controls_satisfied"], s["summary"]["controls_total"]] for s in scans], [28, 18, 22, 8, 10, 9, 8, 9, 8, 18, 15])
    ids = {s["id"] for s in scans}
    names = {s["id"]: s["repository_name"] for s in scans}
    fs = sorted(findings_for(tenant, ids), key=lambda f: -SEV_ORDER[f["severity"]])
    sheet(wb.create_sheet("Findings"), ["Repository", "Severity", "Category", "Rule", "Title", "File", "Line", "Status", "Governance", "Comment"],
          [[names[f["scan_id"]], f["severity"], f["category"], f["rule_id"], f["title"], f["file"], f["line"], f.get("classification"),
            f.get("governance_status"), f.get("comment")] for f in fs], [24, 10, 14, 26, 44, 44, 7, 14, 14, 30])
    rows = []
    for s in scans:
        for c in s.get("controls", []):
            refs = "; ".join(f"{r['framework']}: {', '.join(r['refs'])}" for r in c["framework_refs"])
            rows.append([s["repository_name"], c["control_id"], c["title"], c["status"], c["evidence_count"], refs, "; ".join(c["reasons"])])
    sheet(wb.create_sheet("Controls"), ["Repository", "Control", "Title", "Status", "Evidence", "Framework references", "Reason"], rows, [24, 24, 40, 16, 9, 50, 50])
    comp_rows, ai_rows = [], []
    for s in scans:
        inv = db.get("inventory", tenant, s["id"]) or {}
        comp_rows += [[s["repository_name"], c["name"], c.get("version"), c["ecosystem"], c["manifest"]] for c in inv.get("components", [])]
        ai_rows += [[s["repository_name"], "model", m["name"], m["provider"], ", ".join(m["files"][:5])] for m in inv.get("models", [])]
        ai_rows += [[s["repository_name"], c["category"], c["name"], c["provider"], ", ".join(c["files"][:5])] for c in inv.get("ai_components", [])]
    sheet(wb.create_sheet("Components"), ["Repository", "Package", "Version", "Ecosystem", "Manifest"], comp_rows, [24, 36, 14, 12, 40])
    sheet(wb.create_sheet("AI inventory"), ["Repository", "Kind", "Name", "Provider", "Files"], ai_rows, [24, 16, 36, 24, 60])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/scans/{scan_id}/export/xlsx")
def export_xlsx(scan_id: str, user: dict = Depends(auth.current_user)):
    s = scan_or_404(user, scan_id)
    if s["status"] != "completed":
        raise HTTPException(409, f"scan is {s['status']}")
    audit(user, "export.xlsx", scan_id)
    return Response(_workbook([s], user["tenant_id"]), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{s["repository_name"]}-governance.xlsx"'})


@router.get("/reports/excel")
def report_excel(user: dict = Depends(auth.current_user)):
    scans = latest_scans(user["tenant_id"])
    if not scans:
        raise HTTPException(404, "no completed scans yet")
    audit(user, "export.xlsx_all", "")
    return Response(_workbook(scans, user["tenant_id"]), media_type=XLSX,
                    headers={"Content-Disposition": 'attachment; filename="governance-report.xlsx"'})
